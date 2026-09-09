"""Import d'une liste de locaux depuis un fichier Excel (.xlsx).

Colonnes reconnues (en-têtes insensibles à la casse/aux accents, ordre
libre) : Nom/Local, Type, Surface (m2), Hauteur (m), Fluide, Charge (kg).
Seules les colonnes Nom et Surface sont requises ; les autres sont
optionnelles et complétées avec des valeurs par défaut raisonnables.
"""
from __future__ import annotations

import io
import unicodedata

from openpyxl import load_workbook

COLUMN_ALIASES = {
    "name": {"nom", "local", "piece", "designation", "salle"},
    "room_type": {"type", "typedelocal", "typelocal"},
    "surface_m2": {"surface", "surfacem2", "surfacem²"},
    "height_m": {"hauteur", "hauteurm", "hauteursousplafond", "hsp"},
    "fluid_code": {"fluide", "fluidefrigorigene", "fluid"},
    "charge_kg": {"charge", "chargekg", "chargedefluide"},
}


def _normalize(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return "".join(ch for ch in s.lower() if ch.isalnum())


def parse_rooms_excel(file_bytes: bytes) -> tuple[list[dict], list[str]]:
    warnings: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Impossible de lire le fichier Excel : {exc}") from exc

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ["Le fichier est vide."]

    column_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        normalized = _normalize(str(cell))
        for field, aliases in COLUMN_ALIASES.items():
            if normalized in aliases:
                column_map[idx] = field
                break

    if "name" not in column_map.values() or "surface_m2" not in column_map.values():
        warnings.append(
            "Colonnes attendues introuvables (au minimum 'Nom' et 'Surface'). "
            "En-têtes détectées : " + ", ".join(str(c) for c in header_row if c)
        )
        return [], warnings

    rooms: list[dict] = []
    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None for v in row):
            continue
        entry: dict = {}
        for idx, field in column_map.items():
            if idx < len(row):
                entry[field] = row[idx]

        name = entry.get("name")
        surface = entry.get("surface_m2")
        if name is None or surface is None:
            warnings.append(f"Ligne {row_num} ignorée (nom ou surface manquant).")
            continue
        try:
            surface_val = float(surface)
        except (TypeError, ValueError):
            warnings.append(f"Ligne {row_num} ignorée (surface invalide : {surface!r}).")
            continue

        height_val = entry.get("height_m")
        try:
            height_val = float(height_val) if height_val is not None else 2.5
        except (TypeError, ValueError):
            height_val = 2.5

        charge_val = entry.get("charge_kg")
        try:
            charge_val = float(charge_val) if charge_val is not None else None
        except (TypeError, ValueError):
            charge_val = None

        rooms.append(
            {
                "room_name": str(name).strip(),
                "room_type": str(entry.get("room_type") or "bureau").strip().lower() or "bureau",
                "surface_m2": surface_val,
                "height_m": height_val,
                "fluid_code": (str(entry["fluid_code"]).strip().upper() if entry.get("fluid_code") else None),
                "charge_kg": charge_val,
            }
        )

    if not rooms:
        warnings.append("Aucune ligne exploitable trouvée dans le fichier.")

    return rooms, warnings
