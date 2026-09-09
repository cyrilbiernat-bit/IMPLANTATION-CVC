"""Génération de rapport DOCX."""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt


def build_report_docx(project: dict, results: list[dict], company: dict | None = None) -> bytes:
    doc = Document()

    doc.add_heading("Rapport de vérification NF EN 378-1", level=0)
    doc.add_paragraph("Concentration de fluide frigorigène")

    if company and company.get("company_name"):
        doc.add_paragraph(f"Bureau d'études : {company['company_name']}")

    doc.add_heading("1. Données du projet", level=1)
    doc.add_paragraph(f"Projet : {project.get('name', '-')}")
    if project.get("client"):
        doc.add_paragraph(f"Client : {project['client']}")
    if project.get("building_type"):
        doc.add_paragraph(f"Type de bâtiment : {project['building_type']}")

    doc.add_heading("2. Résumé des hypothèses", level=1)
    doc.add_paragraph(
        "Calculs réalisés selon la méthodologie NF EN 378-1 : Volume = Surface x Hauteur ; "
        "Concentration = Masse relâchée / Volume ; comparaison à la limite applicable "
        "(RCL, LFL pratique, ATEL, ODL)."
    )

    doc.add_heading("3. Résultats de calcul", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "LightGrid-Accent1"
    hdr = table.rows[0].cells
    headers = ["Local", "Fluide", "Charge (kg)", "Volume (m3)", "Concentration (kg/m3)", "Limite", "Conformité"]
    for i, h in enumerate(headers):
        hdr[i].text = h

    for r in results:
        row = table.add_row().cells
        row[0].text = str(r.get("room_name", "-"))
        row[1].text = str(r.get("fluid_code", "-"))
        row[2].text = f"{r.get('charge_kg', 0):.2f}"
        row[3].text = f"{r.get('volume_m3', 0):.2f}"
        row[4].text = f"{r.get('concentration_kg_m3', 0):.4f}"
        row[5].text = f"{r.get('limit_used_kg_m3', 0):.4f} ({r.get('limit_type', '-')})"
        row[6].text = str(r.get("conformity", "-"))

    doc.add_heading("4. Recommandations", level=1)
    for r in results:
        recs = r.get("recommendations") or []
        if not recs:
            continue
        doc.add_paragraph(str(r.get("room_name", "Local")), style="Heading2")
        for rec in recs:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{rec.get('measure')} — {rec.get('reason')} (priorité : {rec.get('priority')})")

    doc.add_heading("5. Références normatives", level=1)
    doc.add_paragraph(
        "NF EN 378-1 — Systèmes de réfrigération et pompes à chaleur. Classification de "
        "sécurité : ISO 817 / ASHRAE 34."
    )
    note = doc.add_paragraph()
    run = note.add_run(
        "Ce rapport est généré automatiquement à titre d'aide à la vérification et doit "
        "être validé par un professionnel qualifié."
    )
    run.italic = True
    run.font.size = Pt(9)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
