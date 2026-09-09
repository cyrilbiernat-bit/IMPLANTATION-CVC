"""Génération de rapport XLSX."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

CONFORMITY_FILLS = {
    "Conforme": PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid"),
    "Conforme sous conditions": PatternFill(start_color="FFE0B2", end_color="FFE0B2", fill_type="solid"),
    "Non conforme": PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid"),
}


def build_report_xlsx(project: dict, results: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats NF EN 378-1"

    ws["A1"] = "Rapport de vérification NF EN 378-1"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = f"Projet : {project.get('name', '-')}"
    ws["A3"] = f"Type de bâtiment : {project.get('building_type', '-')}"

    headers = ["Local", "Type", "Fluide", "Charge (kg)", "Volume (m3)", "Concentration (kg/m3)",
               "Limite (kg/m3)", "Type de limite", "Conformité"]
    header_row = 5
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(results, start=header_row + 1):
        ws.cell(row=i, column=1, value=r.get("room_name", "-"))
        ws.cell(row=i, column=2, value=r.get("room_type", "-"))
        ws.cell(row=i, column=3, value=r.get("fluid_code", "-"))
        ws.cell(row=i, column=4, value=r.get("charge_kg", 0))
        ws.cell(row=i, column=5, value=r.get("volume_m3", 0))
        ws.cell(row=i, column=6, value=r.get("concentration_kg_m3", 0))
        ws.cell(row=i, column=7, value=r.get("limit_used_kg_m3", 0))
        ws.cell(row=i, column=8, value=r.get("limit_type", "-"))
        conf_cell = ws.cell(row=i, column=9, value=r.get("conformity", "-"))
        fill = CONFORMITY_FILLS.get(r.get("conformity"))
        if fill:
            conf_cell.fill = fill

    for col_letter, width in zip("ABCDEFGHI", [24, 16, 10, 12, 12, 20, 14, 14, 22]):
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
